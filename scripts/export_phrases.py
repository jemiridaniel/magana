#!/usr/bin/env python3
"""The app's fixed phrases, for native speakers to record (docs/recording/RECORDING_GUIDE.md).

    backend/.venv/bin/python scripts/export_phrases.py

Writes:
- app/assets/voice/phrases.json: for each language, each phrase's text -> its ID, and the health
  templates' fixed parts, so the app finds the clip for what it is about to say (app/lib/clips.dart).
- docs/recording/magana-phrases.xlsx: one sheet per language (ID, file name, text to record,
  English meaning, kind, text status), plus a "Read me" sheet; and a CSV per language next to it.

Fixed phrases are the app's own texts (app/lib/l10n/app_<lang>.arb) with no blank in them. A
text with a blank ("Line {n} of {total}") is left to the phone's voice: its numbers change.
It also holds what English guidance needs (app/lib/guidance.dart): every phrase's English text by
ID ("_english"), and each phrase with blanks as a pattern per language ("_ui_patterns"), so
Magana's own words can be said in English when a language has no voice yet. Plural and select
texts are left out.
Health templates (content/templates/*.json) are recorded in parts: the words between the number
slots. The numbers themselves are always said by the phone, exactly as printed (CLAUDE.md, rule 1).
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LANGS = {"en": "English (Nigerian accent)", "ha": "Hausa", "yo": "Yorùbá", "ig": "Igbo"}
SLOT = re.compile(r"\{[a-z]\}")


def arb(lang: str) -> dict[str, str]:
    data = json.loads((ROOT / f"app/lib/l10n/app_{lang}.arb").read_text(encoding="utf-8"))
    return {k: v for k, v in data.items() if not k.startswith("@") and isinstance(v, str)}


def templates() -> list[dict]:
    """Each template's parts per language: [(part ID, text)], numbers left out."""
    out = []
    for f in sorted((ROOT / "content/templates").glob("*.json")):
        doc = json.loads(f.read_text(encoding="utf-8"))
        for t in doc.get("templates", []):
            for lang, tr in (t.get("translations") or {}).items():
                text = tr.get("text", "")
                pieces = SLOT.split(text)
                slots = SLOT.findall(text)
                parts, pattern = [], []
                for i, piece in enumerate(pieces, 1):
                    if piece.strip(" .,;:"):
                        pid = f"tpl-{f.stem}-{t['id']}-{i}"
                        parts.append((pid, piece.strip()))
                        pattern.append({"text": piece, "id": pid})
                    elif piece:
                        pattern.append({"text": piece})
                    if i <= len(slots):
                        pattern.append({"slot": slots[i - 1][1]})
                out.append({
                    "id": f"tpl-{f.stem}-{t['id']}", "lang": lang, "text": text, "parts": parts,
                    "english": (t.get("en") or [""])[0], "status": tr.get("status", "draft"),
                    # The whole sentence: its fixed parts (with their clip IDs) and number slots, in order.
                    "pattern": pattern,
                })
    return out


def main() -> int:
    english = arb("en")
    phrases: dict[str, dict[str, str]] = {}
    rows: dict[str, list[list[str]]] = {lang: [] for lang in LANGS}
    for lang in LANGS:
        texts = arb(lang)
        phrases[lang] = {}
        also: dict[str, list[str]] = {}
        for key, text in texts.items():
            if "{" in text or not text.strip():
                continue
            if text in phrases[lang]:  # the same words under another ID: recorded once
                also.setdefault(phrases[lang][text], []).append(key)
                continue
            phrases[lang][text] = key
            rows[lang].append([key, f"{lang}-{key}.m4a", text, english.get(key, ""), "App phrase",
                               "draft" if lang != "en" else "final"])
        for row in rows[lang]:
            if row[0] in also:
                row[4] += " (also used as " + ", ".join(also[row[0]]) + ")"
    patterns: dict[str, list[dict]] = {}
    for t in templates():
        if t["lang"] not in LANGS:
            continue
        patterns.setdefault(t["lang"], []).append({"id": t["id"], "pattern": t["pattern"]})
        n = len(t["parts"])
        for i, (pid, text) in enumerate(t["parts"], 1):
            rows[t["lang"]].append([
                pid, f"{t['lang']}-{pid}.m4a", text, t["english"],
                f"Health template, part {i} of {n} (the phone says the numbers between parts)",
                t["status"]])
    # English guidance: each phrase's English text, and the phrases with blanks as patterns.
    english = {k: v for k, v in arb("en").items() if not re.search(r"\{\w+,\s*(plural|select)", v)}
    ui_patterns: dict[str, list[dict]] = {}
    for lang in LANGS:
        if lang == "en":
            continue
        for key, text in arb(lang).items():
            if "{" not in text or key not in english or re.search(r"\{\w+,\s*(plural|select)", text):
                continue
            pieces = [({"slot": p[1:-1]} if re.fullmatch(r"\{\w+\}", p) else {"text": p})
                      for p in re.split(r"(\{\w+\})", text) if p]
            ui_patterns.setdefault(lang, []).append({"id": key, "pattern": pieces})
    out = ROOT / "app/assets/voice/phrases.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"_about": "Made by scripts/export_phrases.py; do not edit.",
                               **phrases, "_patterns": patterns, "_english": english,
                               "_ui_patterns": ui_patterns},
                              ensure_ascii=False, indent=1, sort_keys=False) + "\n", encoding="utf-8")
    clips = ROOT / "app/assets/voice/clips.json"
    if not clips.exists():
        clips.write_text('{"_about": "Made by scripts/import_clips.py from the recordings."}\n', encoding="utf-8")

    docs = ROOT / "docs/recording"
    docs.mkdir(parents=True, exist_ok=True)
    header = ["Phrase ID", "File name", "Text to record", "English meaning", "Kind", "Text status",
              "Recorded? (yes/no)", "Notes"]
    for lang, lines in rows.items():
        with open(docs / f"phrases-{lang}.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(lines)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl isn't installed: CSV files only", file=sys.stderr)
    else:
        wb = Workbook()
        about = wb.active
        about.title = "Read me"
        for line in [
            "Magana: phrases to record",
            "",
            "One sheet per language. Record each 'Text to record' exactly as written, one file per row,",
            "named as in 'File name' (for example ha-homeTitle.m4a). Then mark 'Recorded?' yes.",
            "How to record: docs/recording/RECORDING_GUIDE.md (quiet room, phone voice memo, file names).",
            "",
            "Numbers are never recorded: the phone always says them exactly as printed.",
            "Health template parts are the words between the numbers.",
            "'Text status': draft texts may still change; tell us if a text reads wrong (Notes column).",
            "Hausa, Yorùbá and Igbo texts are drafts until a native speaker has checked them.",
        ]:
            about.append([line])
        about["A1"].font = Font(bold=True, size=14)
        about.column_dimensions["A"].width = 100
        for lang, lines in rows.items():
            ws = wb.create_sheet(LANGS[lang][:31])
            ws.append(header)
            for line in lines:
                ws.append(line + ["", ""])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0A6B4B")
            ws.freeze_panes = "A2"
            for col, width in zip("ABCDEFGH", [34, 40, 60, 50, 36, 18, 18, 30]):
                ws.column_dimensions[col].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"G2:G{len(lines) + 1}")
        wb.save(docs / "magana-phrases.xlsx")
    total = {lang: len(lines) for lang, lines in rows.items()}
    print(f"phrases: {total}; app/assets/voice/phrases.json and docs/recording/ written")
    return 0


if __name__ == "__main__":
    sys.exit(main())
