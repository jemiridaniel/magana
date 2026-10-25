"""Build docs/review/: files non-developers can open on a phone to check Magana's content.

    backend/.venv/bin/python scripts/build_review_pack.py

Needs openpyxl (pip install openpyxl), pandoc and Google Chrome (for the PDF), and the
iba-natlas repo next to this one (for the danger-sign questions). Re-run after any change to
the templates, sample questions, word lists, app strings or privacy notice.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "review"
IBA = ROOT.parent / "iba-natlas" / "app"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
LANGS = {"ha": "Hausa", "yo": "Yorùbá", "ig": "Igbo"}
HEAD = PatternFill("solid", fgColor="0A6B4B")
CHECK = "Correct? (yes / fix)"


def sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF", size=12)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def finish(ws) -> None:
    """Wrap every cell, and a yes / fix drop-down in the Correct? column."""
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(size=12, bold=bool(c.font and c.font.bold))
    headers = [c.value for c in ws[1]]
    if CHECK in headers:
        col = ws.cell(1, headers.index(CHECK) + 1).column_letter
        dv = DataValidation(type="list", formula1='"yes,fix"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(ws.max_row, 2)}")


def how_to(wb: Workbook, lines: list[str]) -> None:
    ws = wb.active
    ws.title = "How to review"
    ws.column_dimensions["A"].width = 90
    for i, line in enumerate(lines, 1):
        ws.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).font = Font(size=13, bold=i == 1)


# ---------- 1. medicine-label templates

def templates() -> Path:
    doc = json.loads((ROOT / "content" / "templates" / "medicine_label.json").read_text())
    wb = Workbook()
    how_to(wb, [
        "Medicine-label phrase templates: review",
        "Magana reads these phrases in Hausa, Yorùbá or Igbo instead of English, only once a "
        "native speaker has checked them and Dr. Amos has approved them clinically. Until then "
        "every medicine-label line is read in English.",
        "{a}, {b}, {c}, {d} are where Magana puts the number exactly as printed on the label "
        "(for example '1', '24', '08/2027'). Please don't change or move them unless the word "
        "order needs it; every translation must keep the same letters.",
        "In the 'Correct?' column choose yes, or fix and write the correct wording in Notes.",
        "Several English forms on one row mean the same phrase printed in different ways.",
        "Source: content/templates/medicine_label.json in the Magana repo.",
    ])
    ws = sheet(wb, "Templates", ["#", "English", "Hausa", "Yorùbá", "Igbo", "Status", CHECK, "Notes", "ID"],
               [4, 34, 34, 34, 34, 14, 14, 34, 22])
    for i, t in enumerate(doc["templates"], 1):
        tr = t["translations"]
        statuses = {tr[l]["status"] for l in LANGS}
        ws.append([i, "\n".join(t["en"]), tr["ha"]["text"], tr["yo"]["text"], tr["ig"]["text"],
                   statuses.pop() if len(statuses) == 1 else ", ".join(f"{l}: {tr[l]['status']}" for l in LANGS),
                   "", "", t["id"]])
    finish(ws)
    path = OUT / "medicine_label_templates.xlsx"
    wb.save(path)
    return path


# ---------- 2. sample questions

BANDS = {"primary_1_3": "Primary 1–3", "primary_4_6": "Primary 4–6", "jss": "JSS 1–3", "sss": "SSS 1–3"}
SUBJECTS = {"english": "English", "mathematics": "Mathematics", "basic_science": "Basic Science", "biology": "Biology"}
BAND_ORDER = list(BANDS)


def sample_questions() -> Path:
    packs = []
    for p in (ROOT / "app" / "assets" / "quiz").glob("*.json"):
        packs.append(json.loads(p.read_text()))
    packs.sort(key=lambda d: (BAND_ORDER.index(d["band"]), list(SUBJECTS).index(d["subject"])))
    wb = Workbook()
    how_to(wb, [
        "Sample quiz questions: review by a teacher",
        "These are Magana's sample practice questions, written by the Magana team in our own "
        "words (not past questions). Learners hear each question and its options read aloud and "
        "answer by voice or by tapping.",
        "Please check: is the question right for the level? Is the right answer really right, "
        "and the others clearly wrong? Is it clear when heard, not read? Primary questions must "
        "have 2 or 3 options and 15 words or fewer.",
        "In the 'Correct?' column choose yes, or fix and write what to change in Notes.",
        "'Said after the answer' is what Magana says once the learner has answered.",
        f"{sum(len(d['questions']) for d in packs)} questions in {len(packs)} packs. Source: app/assets/quiz/.",
    ])
    ws = sheet(wb, "Questions", ["#", "Level", "Subject", "Question", "Options", "Right answer",
                                 CHECK, "Notes", "Said after the answer", "ID"],
               [5, 13, 14, 40, 34, 24, 14, 34, 34, 26])
    n = 0
    for d in packs:
        for q in d["questions"]:
            n += 1
            letters = "ABCD"
            ws.append([n, BANDS[d["band"]], SUBJECTS[d["subject"]], q["question"],
                       "\n".join(f"{letters[i]}. {o}" for i, o in enumerate(q["options"])),
                       f"{letters[q['answer']]}. {q['options'][q['answer']]}", "", "",
                       q.get("explanation", ""), q["id"]])
    finish(ws)
    path = OUT / "sample_questions.xlsx"
    wb.save(path)
    return path


# ---------- 3. language review: number words, answer words, spoken lines

SECTIONS = [  # (where the learner hears it, key prefixes), first match wins
    ("Server messages (readings, warnings, health, money)", ("msg_",)),
    ("Getting started", ("chooseLanguage", "welcome", "age", "privacyNote", "ready", "startLabel", "consent",
                         "classroom", "class", "startMagana", "hearAgain")),
    ("Practice quiz", ("quiz",)),
    ("Explain", ("explain", "readExactly", "readAgain")),
    ("Reading a page", ("readPage", "readingPage", "readFailed", "readAnyPaper", "line", "unclear", "endOfPage",
                        "where", "table")),
    ("Talk to a health worker", ("talkToNurse", "quickCheck", "callback")),
    ("Voice commands", ("command", "voice")),
]


def section_of(key: str) -> str | None:
    for name, prefixes in SECTIONS:
        if key.startswith(prefixes):
            return name
    return None


def language_review() -> Path:
    arb = {l: json.loads((ROOT / "app" / "lib" / "l10n" / f"app_{l}.arb").read_text()) for l in ("en", *LANGS)}
    sets = json.loads((ROOT / "content" / "answers" / "closed_sets.json").read_text())
    wb = Workbook()
    how_to(wb, [
        "Language review: Hausa, Yorùbá, Igbo",
        "One tab per language. Each tab has three parts:",
        "1. Number words: words Magana must never let an AI write by itself (numbers and units). "
        "Add missing spellings and variants in Notes. 'Only with tone marks' means the word is "
        "checked only when written with its marks, because without them it is another word. "
        "'Not checked' words are also everyday words (like 'one' / 'a'): tell us if any can be "
        "checked safely.",
        "2. Quiz answer words: how learners say A, B, C, D, yes, no, true and false. Add other "
        "ways people say them.",
        "3. Spoken lines: what the app says, in English and in your language. Check that it is "
        "correct, natural and easy to understand when heard.",
        "Words like {name} or {letters} are filled in by the app: keep them as they are.",
        "In 'Correct?' choose yes, or fix and write the correct wording in Notes.",
    ])
    for lang, name in LANGS.items():
        ws = sheet(wb, name, ["Part", "What", "English", name, CHECK, "Notes"], [16, 26, 40, 40, 14, 34])
        words = json.loads((ROOT / "content" / "number_words" / f"{lang}.json").read_text())
        for kind, key in (("Number word", "words"), ("Number word (only with tone marks)", "exact_only"),
                          ("Not checked (also an everyday word)", "ambiguous_not_checked"), ("Unit word", "units")):
            for w in words[key]:
                ws.append(["Number words", kind, "", w, "", ""])
        for kind, label in (("letters", "Letter"), ("yes_no", "Yes / no"), ("true_false", "True / false")):
            for answer, forms in sets[kind].get(lang, {}).items():
                ws.append(["Quiz answer words", f"{label}: {answer}", ", ".join(sets[kind]["en"][answer]),
                           ", ".join(forms), "", ""])
        for key, english in arb["en"].items():
            if key.startswith("@"):
                continue
            where = section_of(key)
            if where and key in arb[lang]:
                ws.append(["Spoken lines", where, english, arb[lang][key], "", ""])
        finish(ws)
    path = OUT / "language_review.xlsx"
    wb.save(path)
    return path


# ---------- 4. privacy notice as PDF

CSS = """body{font-family:-apple-system,Helvetica,Arial,sans-serif;font-size:17px;line-height:1.5;
max-width:46em;margin:2em auto;padding:0 1em;color:#14171A}h1{color:#0A6B4B}
table{border-collapse:collapse;width:100%;font-size:15px}td,th{border:1px solid #8a8f94;padding:6px;
vertical-align:top}th{background:#E3F1EA}blockquote{background:#FFF1D6;margin:0;padding:.5em 1em}"""


def privacy_pdf() -> Path:
    path = OUT / "PRIVACY.pdf"
    with tempfile.TemporaryDirectory() as tmp:
        html = Path(tmp) / "privacy.html"
        css = Path(tmp) / "style.css"
        css.write_text(CSS)
        subprocess.run(["pandoc", str(ROOT / "docs" / "PRIVACY.md"), "-s", "-o", str(html), "--css", str(css),
                        "--metadata", "pagetitle=Magana privacy notice (draft)", "--embed-resources"], check=True)
        subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={path}", html.as_uri()], check=True, capture_output=True)
    return path


# ---------- 5. Ibà self-report questions, exact English wording

def danger_signs() -> Path:
    sys.path.insert(0, str(IBA))
    from backend.app.questions import questions  # noqa: E402  (iba-natlas, read only)

    qs = questions("self_report", "en")
    magana = json.loads((ROOT / "app" / "lib" / "l10n" / "app_en.arb").read_text())
    sight = [k for k in magana if re.fullmatch(r"msg_q_\w+_no_sight", k)]
    lines = [
        "Ibà self-report questions (the public, through Magana): exact English wording",
        "Source: iba-natlas app/docs/CLINICAL_REVIEW.md item 7, wording in app/backend/app/questions.py",
        "Status: pending clinical sign-off by Dr. Oluba Amos",
        "",
    ]
    for i, q in enumerate(qs, 1):
        when = f" (asked if the answer to {q['ask_if']} is yes; answer: number of days)" if q.get("ask_if") else ""
        lines.append(f"{i}. {q['text']}{when}")
    lines += [
        "",
        "Rules (item 7): any of questions 3 to 12 answered yes -> emergency: go to the nearest clinic now, "
        "or call 112. Fever for 2 days or more with no danger sign -> urgent: a nurse calls back within "
        "2 hours. Otherwise -> routine: a nurse calls back during working hours.",
        "",
        "Magana's versions for blind users, asked instead of the sight-based questions (DRAFT, also for "
        "Dr. Amos's sign-off; Magana app/lib/l10n/app_en.arb):",
    ]
    for i, k in enumerate(sight, 1):
        lines.append(f"{i}. ({k.removeprefix('msg_q_').removesuffix('_no_sight')}) {magana[k]}")
    lines += ["", "Every question in Magana also offers \"I can't tell\": never counted as no; on a danger sign "
              f"it raises priority to at least urgent and Magana says: \"{magana['msg_ask_someone_to_check']}\""]
    path = OUT / "danger_signs.txt"
    path.write_text("\n".join(lines) + "\n")
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for build in (templates, sample_questions, language_review, privacy_pdf, danger_signs):
        print(build())
